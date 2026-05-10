from __future__ import annotations

from contextlib import contextmanager
from pathlib import Path
from tempfile import TemporaryDirectory

import pytest

from graph_rag_app.corpus import chunk_text, load_corpus_text
from graph_rag_app.indexing import build_index, load_index, _default_embedding_client
from graph_rag_app.retrieval import SearchResult


def _embedding_available() -> bool:
    """Check if the embedding API is reachable."""
    client = _default_embedding_client()
    if client is None:
        return False
    try:
        client.embed_query("connectivity test")
        return True
    except Exception:
        return False


@contextmanager
def built_markdown_index():
    with TemporaryDirectory() as temp_dir:
        root = Path(temp_dir)
        kb_path = root / "kb"
        output_dir = root / "index"
        kb_path.mkdir()

        (kb_path / "agent_orchestration.md").write_text(
            "# Agent Orchestration\n\n"
            "AI agents coordinate tools, memory, and planning loops for reliable work.\n",
            encoding="utf-8",
        )
        (kb_path / "retrieval_agents.md").write_text(
            "# Retrieval Agents\n\n"
            "Retrieval augmented agents search documents before answering questions.\n",
            encoding="utf-8",
        )
        (kb_path / "evaluation.md").write_text(
            "# Evaluation\n\n"
            "Agent evaluation checks grounded responses and task completion quality.\n",
            encoding="utf-8",
        )

        corpus_text = load_corpus_text(kb_path)
        assert chunk_text(corpus_text)

        build_index(kb_path, output_dir)
        yield kb_path, output_dir


def test_build_and_query_hybrid():
    if not _embedding_available():
        pytest.skip("Embedding API is not available")
    with built_markdown_index() as (_, index_dir):
        retriever = load_index(index_dir)
        try:
            results = retriever.retrieve("agent orchestration", top_k=3)
        finally:
            retriever.close()

    assert results
    assert all(isinstance(result, SearchResult) for result in results)
    assert results[0].score > 0


def test_build_and_query_dense():
    if not _embedding_available():
        pytest.skip("Embedding API is not available")
    with built_markdown_index() as (_, index_dir):
        retriever = load_index(index_dir)
        try:
            results = retriever.retrieve("retrieval agents", top_k=3, strategy="dense")
        finally:
            retriever.close()

    assert results
    assert all(isinstance(result, SearchResult) for result in results)


def test_build_and_query_sparse():
    with built_markdown_index() as (_, index_dir):
        retriever = load_index(index_dir)
        try:
            results = retriever.retrieve("evaluation grounded", top_k=3, strategy="sparse")
        finally:
            retriever.close()

    assert results
    assert all(isinstance(result, SearchResult) for result in results)


def test_load_index_close():
    with built_markdown_index() as (_, index_dir):
        retriever = load_index(index_dir)
        retriever.close()


def test_build_from_docx():
    try:
        from docx import Document
    except ImportError:
        pytest.skip("python-docx is not available")

    with TemporaryDirectory() as temp_dir:
        root = Path(temp_dir)
        docx_path = root / "agent_notes.docx"
        output_dir = root / "index"

        document = Document()
        document.add_heading("AI Agent Notes", level=1)
        document.add_paragraph(
            "AI agents use retrieval, planning, and tool execution to answer grounded questions."
        )
        document.save(docx_path)

        build_index(docx_path, output_dir)
        retriever = load_index(output_dir)
        try:
            results = retriever.retrieve("tool execution", top_k=3)
        finally:
            retriever.close()

    assert results
    assert all(isinstance(result, SearchResult) for result in results)
